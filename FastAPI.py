from fastapi import FastAPI, UploadFile, File
from pptx import Presentation
import shutil
import os

app = FastAPI()

UPLOAD_DIR = "uploads"
os.makedirs(UPLOAD_DIR, exist_ok=True)

@app.post("/upload-ppt")
async def upload_ppt(file: UploadFile = File(...)):
    if not file.filename.endswith(".pptx"):
        return {"error": "Only .pptx files are allowed"}

    file_path = os.path.join(UPLOAD_DIR, file.filename)

    with open(file_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)

    prs = Presentation(file_path)

    slides = []

    for index, slide in enumerate(prs.slides, start=1):
        text = ""
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text += shape.text + "\n"

        slides.append({
            "slide_number": index,
            "text": text.strip()
        })

    return {
        "filename": file.filename,
        "total_slides": len(slides),
        "slides": slides
    }
@app.get("/")
async def root():   
    return {"message": "Welcome to the PPTX Upload API. Use the /upload-ppt endpoint to upload a .pptx file."}          



